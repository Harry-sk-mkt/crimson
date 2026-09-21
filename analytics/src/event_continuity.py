"""이벤트 간 연속 참석(연속성) 분석: 저번 이벤트도 보고 이번 이벤트도 보는 단골.

WHY: 단골(regular)은 '한 해에 몇 번 왔나'보다 '이어서 오는가'로 더 잘 드러난다. 연속 참석률이 떨어졌다면
이벤트가 있어도 이어 보는 습관이 약해졌다는 뜻이다. 멤버 리포트에는 이벤트 날짜가 없어서, 날짜가 있는
사용자 제공 이벤트 일정표(2025 8개, 2026 6개)와 캠페인 이름으로 매칭한 이벤트만 쓴다.
주의: 이벤트 간격이 연도별로 다르고(2025는 3~14일, 2026은 4~7일) 이벤트 수가 적다.
"""
import datetime as dt
import math

import openpyxl
import pandas as pd

import config
from load_campaign_members import load_all


def load_calendar(path=None) -> pd.DataFrame:
    """일정표 -> 컬럼: date, campaign. 헤더 라벨 위치를 찾아 그 아래 (날짜, WB- 이름) 행을 읽는다."""
    ws = openpyxl.load_workbook(path or config.EVENT_CALENDAR_FILE, data_only=True)[config.EVENT_CALENDAR_SHEET]
    heads = [c for row in ws.iter_rows() for c in row if c.value == config.EVENT_DATE_HEADER]
    rows = []
    for h in heads:
        col = h.column
        for r in range(h.row + 1, ws.max_row + 1):
            d, n = ws.cell(r, col).value, ws.cell(r, col + 1).value
            if isinstance(d, (dt.datetime, dt.date)) and isinstance(n, str) and n.strip().startswith(config.CAMPAIGN_PREFIX):
                rows.append({"date": d.date() if isinstance(d, dt.datetime) else d, "campaign": n.strip()})
    return pd.DataFrame(rows).drop_duplicates().sort_values("date").reset_index(drop=True)


def build_slots(cal: pd.DataFrame, members: pd.DataFrame, year: int, merge: bool = True):
    """연도별 이벤트 슬롯 목록: [(date, campaign, 참석 이메일 집합)] 날짜순.
    제외 캠페인, 멤버가 너무 적은 이벤트를 빼고, merge면 레코딩을 라이브 슬롯에 합친다."""
    m = members.copy()
    m["Campaign Name"] = m["Campaign Name"].str.strip()
    size = m.groupby("Campaign Name").size()
    slots = []
    for _, r in cal.iterrows():
        name = r["campaign"]
        if r["date"].year != year or name in config.EXCLUDED_CAMPAIGNS:
            continue
        if merge and name in config.MERGED_INTO:
            continue                                   # 레코딩 행은 라이브 슬롯에 흡수
        if size.get(name, 0) < config.MIN_EVENT_MEMBERS:
            continue
        names = [name]
        if merge:
            names += [k for k, v in config.MERGED_INTO.items() if v == name]
        att = set(m[(m["Campaign Name"].isin(names)) & (m["Member Status"] == "Attended")]["Email"])
        slots.append((r["date"], name, att))
    return sorted(slots, key=lambda s: s[0])


def pair_table(slots) -> pd.DataFrame:
    """인접한 이벤트 쌍별: 간격(일), 앞 이벤트 참석자 수, 둘 다 참석, 이어 본 비율, 뒤 이벤트 참석자 중 이어온 비중."""
    rows = []
    for (d0, n0, a0), (d1, n1, a1) in zip(slots, slots[1:]):
        both = len(a0 & a1)
        rows.append({"from": d0, "to": d1, "gap_days": (d1 - d0).days, "n_prev": len(a0), "n_next": len(a1), "both": both,
                     "continue_rate": both / len(a0) if a0 else float("nan"),
                     "repeat_share": both / len(a1) if a1 else float("nan")})
    return pd.DataFrame(rows)


def wilson(k: int, n: int, z: float = 1.96):
    if n == 0:
        return float("nan"), float("nan")
    p = k / n
    d = 1 + z * z / n
    c = (p + z * z / (2 * n)) / d
    h = z * math.sqrt(p * (1 - p) / n + z * z / (4 * n * n)) / d
    return c - h, c + h


def pooled(pairs: pd.DataFrame, max_gap=None):
    p = pairs if max_gap is None else pairs[pairs["gap_days"] <= max_gap]
    k, n = int(p["both"].sum()), int(p["n_prev"].sum())
    return {"pairs": len(p), "both": k, "prev": n, "rate": k / n if n else float("nan"), "ci": wilson(k, n)}


def lag_rate(slots, lag: int):
    """이벤트 k 참석자 중 k+lag 이벤트에도 참석한 비율 (모든 k를 합산). lag=1이 '바로 다음'.
    lag가 커져도 비율이 비슷하면 '이어서 오는' 사람이 아니라 '자주 오는' 사람이 있다는 뜻이다."""
    k = n = 0
    for i in range(len(slots) - lag):
        a0, a1 = slots[i][2], slots[i + lag][2]
        k += len(a0 & a1)
        n += len(a0)
    return k, n


def two_prop_z(k1: int, n1: int, k2: int, n2: int):
    """두 비율 차이(2 - 1)의 근사 z와 양측 p. 같은 사람이 여러 쌍에 들어가 독립이 아니므로 참고용이다."""
    p1, p2, p = k1 / n1, k2 / n2, (k1 + k2) / (n1 + n2)
    se = math.sqrt(p * (1 - p) * (1 / n1 + 1 / n2))
    z = (p2 - p1) / se
    return z, math.erfc(abs(z) / math.sqrt(2))


def streaks(slots):
    """사람별 최장 연속 참석 이벤트 수 -> (참석자 수, 연속 2회+ 인원, 연속 3회+ 인원)."""
    people = set().union(*[s[2] for s in slots]) if slots else set()
    best = {}
    for e in people:
        run = mx = 0
        for _, _, a in slots:
            run = run + 1 if e in a else 0
            mx = max(mx, run)
        best[e] = mx
    return len(people), sum(v >= 2 for v in best.values()), sum(v >= 3 for v in best.values())


# ---- 테스트: 기대값과 실제값 비교 ----
def _toy_slots():
    d = lambda i: dt.date(2026, 7, i)
    return [(d(1), "A", {"a", "b", "c"}), (d(8), "B", {"a", "b", "d"}), (d(15), "C", {"a", "e"}), (d(22), "D", {"e"})]


def test_pair_table_toy():
    p = pair_table(_toy_slots())
    assert p["gap_days"].tolist() == [7, 7, 7]
    assert p["both"].tolist() == [2, 1, 1]                 # A~B: a,b / B~C: a / C~D: e
    assert p["n_prev"].tolist() == [3, 3, 2]
    assert abs(p["continue_rate"].iloc[0] - 2 / 3) < 1e-9
    assert abs(p["repeat_share"].iloc[1] - 1 / 2) < 1e-9   # C 참석자 2명 중 B에서 이어온 사람 1명(a)


def test_streaks_toy():
    # a: A,B,C 연속 3 / b: A,B 2 / c: 1 / d: 1 / e: C,D 2
    assert streaks(_toy_slots()) == (5, 3, 1)


def test_pooled_and_wilson():
    p = pooled(pair_table(_toy_slots()))
    assert (p["both"], p["prev"]) == (4, 8) and p["rate"] == 0.5
    lo, hi = p["ci"]
    assert 0.2 < lo < 0.5 < hi < 0.8


def test_lag_rate_toy():
    # lag1: A->B 2/3, B->C 1/3, C->D 1/2 => 4/8 / lag2: A->C 1/3, B->D 0/3 => 1/6
    assert lag_rate(_toy_slots(), 1) == (4, 8)
    assert lag_rate(_toy_slots(), 2) == (1, 6)


def test_two_prop_z():
    z, p = two_prop_z(30, 100, 30, 100)          # 같은 비율이면 z=0, p=1
    assert abs(z) < 1e-9 and abs(p - 1) < 1e-9
    z, p = two_prop_z(10, 100, 40, 100)          # 큰 차이는 유의
    assert z > 4 and p < 0.001


def test_calendar_matches_member_campaigns():
    cal = load_calendar()
    names = set(load_all()["Campaign Name"].str.strip())
    assert len(cal) == 14, len(cal)
    assert set(cal["campaign"]) <= names, set(cal["campaign"]) - names


def test_slots_rules_on_real_data():
    cal, mem = load_calendar(), load_all()
    s25 = build_slots(cal, mem, 2025)
    s26 = build_slots(cal, mem, 2026)
    assert not any(s[1] in config.EXCLUDED_CAMPAIGNS for s in s25)        # 미국 공동 이벤트 제외
    assert len(s25) == 6 and len(s26) == 5, (len(s25), len(s26))          # 8/2(소규모) 제외, 레코딩은 라이브에 합침
    live_name = list(config.MERGED_INTO.values())[0]
    merged = [s for s in s26 if s[1] == live_name][0]
    live_only = [s for s in build_slots(cal, mem, 2026, merge=False) if s[1] == live_name][0]
    assert merged[2] >= live_only[2] and len(merged[2]) > len(live_only[2])   # 합치면 참석자가 늘어난다


if __name__ == "__main__":
    for t in [v for k, v in sorted(globals().items()) if k.startswith("test_")]:
        t()
    print("모든 테스트 통과")
    cal, mem = load_calendar(), load_all()
    base = {y: build_slots(cal, mem, y) for y in (2025, 2026)}
    print("\n=== 비교: 2025 vs 2026 (레코딩 합침) ===")
    for lag in (1, 2, 3):
        (k5, n5), (k6, n6) = lag_rate(base[2025], lag), lag_rate(base[2026], lag)
        z, p = two_prop_z(k5, n5, k6, n6)
        print(f"  {lag}개 뒤 이벤트에도 참석: 2025 {k5}/{n5} ({k5/n5*100:.0f}%) -> 2026 {k6}/{n6} ({k6/n6*100:.0f}%)  z={z:.2f}, 근사 p={p:.3f}")
    r5, r6 = [pair_table(base[y]) for y in (2025, 2026)]
    rep5 = (int(r5['both'].sum()), int(r5['n_next'].sum())); rep6 = (int(r6['both'].sum()), int(r6['n_next'].sum()))
    z, p = two_prop_z(rep5[0], rep5[1], rep6[0], rep6[1])
    print(f"  이번 이벤트 참석자 중 직전 이벤트도 본 사람: 2025 {rep5[0]}/{rep5[1]} ({rep5[0]/rep5[1]*100:.0f}%) -> 2026 {rep6[0]}/{rep6[1]} ({rep6[0]/rep6[1]*100:.0f}%)  z={z:.2f}, 근사 p={p:.3f}")
    for label, merge in (("레코딩을 라이브와 합침(기준)", True), ("레코딩 이벤트를 별도 슬롯으로 (합치지 않음)", False)):
        print(f"\n=== {label} ===")
        for y in (2025, 2026):
            slots = build_slots(cal, mem, y, merge=merge)
            pt = pair_table(slots)
            print(f"\n[{y}] 이벤트 {len(slots)}개: " + ", ".join(f"{s[0]:%m/%d}({len(s[2])}명)" for s in slots))
            for _, r in pt.iterrows():
                print(f"  {r['from']:%m/%d}->{r['to']:%m/%d} ({r['gap_days']:2d}일) 앞 참석 {r['n_prev']:2d}명 중 이어 봄 {r['both']:2d}명 ({r['continue_rate']*100:3.0f}%) | 뒤 참석자 {r['n_next']:2d}명 중 이어온 사람 {r['repeat_share']*100:3.0f}%")
            for g, name in ((None, "모든 인접 쌍"), (config.WEEKLY_MAX_GAP_DAYS, f"간격 {config.WEEKLY_MAX_GAP_DAYS}일 이하 쌍")):
                p = pooled(pt, g)
                print(f"  {name}: {p['both']}/{p['prev']} = {p['rate']*100:.0f}% (95% CI {p['ci'][0]*100:.0f}~{p['ci'][1]*100:.0f}%, 쌍 {p['pairs']}개)")
            n, s2, s3 = streaks(slots)
            print(f"  참석자 {n}명 중 연속 2회+ {s2}명 ({s2/n*100:.0f}%), 연속 3회+ {s3}명 ({s3/n*100:.0f}%)")
